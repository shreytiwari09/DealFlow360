"""Sales reporting (PRD A7).

Pure filtering/aggregation over quotations, plus export rendering. The
filters match PRD A7's own "Reporting Filters" list exactly: Period,
Sales Team/Rep, Approval Status, Product/Category.
"""

from __future__ import annotations

from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal
from io import BytesIO

from fpdf import FPDF
from openpyxl import Workbook
from openpyxl.styles import Font
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy.orm import selectinload

from app.models.quotation import Quotation, QuotationLine


@dataclass(frozen=True)
class ReportFilters:
    date_from: date | None = None
    date_to: date | None = None
    owner_id: int | None = None
    status: str | None = None
    category_id: int | None = None


@dataclass(frozen=True)
class ReportRow:
    quote_number: str
    customer_name: str
    owner_name: str
    status: str
    total_amount: Decimal
    discount_amount: Decimal
    blended_risk_score: Decimal
    created_at: datetime


@dataclass(frozen=True)
class Report:
    rows: tuple[ReportRow, ...]
    total_amount: Decimal
    total_discount: Decimal


async def build_report(
    session: AsyncSession, filters: ReportFilters, *, owner_scope: int | None
) -> Report:
    """`owner_scope` is the caller's own ownership restriction (a rep may
    only ever report on their own deals, same as every other list in this
    app) — kept separate from `filters.owner_id`, which is the caller
    deliberately narrowing a report they're already allowed to see broadly
    (a manager filtering "just this rep's numbers")."""
    statement = select(Quotation).options(
        selectinload(Quotation.customer), selectinload(Quotation.owner)
    )

    if owner_scope is not None:
        statement = statement.where(Quotation.owner_id == owner_scope)
    if filters.owner_id is not None:
        statement = statement.where(Quotation.owner_id == filters.owner_id)
    if filters.status is not None:
        statement = statement.where(Quotation.status == filters.status)
    if filters.date_from is not None:
        statement = statement.where(Quotation.created_at >= filters.date_from)
    if filters.date_to is not None:
        statement = statement.where(Quotation.created_at < filters.date_to)
    if filters.category_id is not None:
        subquery = (
            select(QuotationLine.quotation_id)
            .join(QuotationLine.product)
            .where(QuotationLine.product.has(category_id=filters.category_id))
        )
        statement = statement.where(Quotation.id.in_(subquery))

    statement = statement.order_by(Quotation.created_at.desc())
    quotations = (await session.execute(statement)).scalars().all()

    rows = tuple(
        ReportRow(
            quote_number=q.quote_number,
            customer_name=q.customer.name,
            owner_name=q.owner.full_name,
            status=q.status,
            total_amount=q.total_amount,
            discount_amount=q.discount_amount,
            blended_risk_score=q.blended_risk_score,
            created_at=q.created_at,
        )
        for q in quotations
    )
    return Report(
        rows=rows,
        total_amount=sum((r.total_amount for r in rows), Decimal("0")),
        total_discount=sum((r.discount_amount for r in rows), Decimal("0")),
    )


# --- Export rendering ---------------------------------------------------------

_COLUMNS = ("Quote #", "Customer", "Owner", "Status", "Total", "Discount", "Risk Score", "Created")


def _row_values(row: ReportRow) -> tuple[str, ...]:
    return (
        row.quote_number,
        row.customer_name,
        row.owner_name,
        row.status,
        f"{row.total_amount:.2f}",
        f"{row.discount_amount:.2f}",
        f"{row.blended_risk_score:.2f}",
        row.created_at.strftime("%Y-%m-%d"),
    )


def render_xlsx(report: Report) -> bytes:
    """openpyxl — pure Python, no system dependency (Section 0.5 checkpoint,
    see PROJECT_CONTEXT.md Important Technical Decisions)."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Quotations"

    for col, header in enumerate(_COLUMNS, start=1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = Font(bold=True)

    for row_index, row in enumerate(report.rows, start=2):
        for col, value in enumerate(_row_values(row), start=1):
            ws.cell(row=row_index, column=col, value=value)

    total_row = len(report.rows) + 2
    ws.cell(row=total_row, column=1, value="TOTAL").font = Font(bold=True)
    ws.cell(row=total_row, column=5, value=f"{report.total_amount:.2f}").font = Font(bold=True)
    ws.cell(row=total_row, column=6, value=f"{report.total_discount:.2f}").font = Font(bold=True)

    for col in range(1, len(_COLUMNS) + 1):
        ws.column_dimensions[chr(64 + col)].width = 16

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def render_pdf(report: Report) -> bytes:
    """fpdf2 — basic tabular styling only (no charts), which is exactly what
    a sales report needs (Section 0.5 checkpoint)."""
    pdf = FPDF(orientation="L", unit="mm", format="A4")
    pdf.add_page()
    pdf.set_font("Helvetica", "B", 14)
    pdf.cell(0, 10, "DealFlow360 - Sales Report", new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("Helvetica", "", 9)

    widths = (25, 45, 35, 28, 25, 25, 22, 25)
    pdf.set_font("Helvetica", "B", 9)
    for header, width in zip(_COLUMNS, widths, strict=True):
        pdf.cell(width, 8, header, border=1)
    pdf.ln()

    pdf.set_font("Helvetica", "", 9)
    for row in report.rows:
        for value, width in zip(_row_values(row), widths, strict=True):
            pdf.cell(width, 7, value, border=1)
        pdf.ln()

    pdf.set_font("Helvetica", "B", 9)
    pdf.cell(sum(widths[:4]), 8, "TOTAL", border=1)
    pdf.cell(widths[4], 8, f"{report.total_amount:.2f}", border=1)
    pdf.cell(widths[5], 8, f"{report.total_discount:.2f}", border=1)
    pdf.cell(widths[6] + widths[7], 8, "", border=1)

    return bytes(pdf.output())
