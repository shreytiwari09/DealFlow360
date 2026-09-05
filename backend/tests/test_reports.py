"""Reporting tests (PRD A7).

`build_report` needs the real schema (it filters/joins real quotations), so
it is tested against the database. `render_xlsx`/`render_pdf` are pure
rendering over a plain `Report` dataclass - tested without a database at all,
just checking the export is well-formed and non-empty.
"""

import itertools
from datetime import UTC, datetime
from decimal import Decimal
from io import BytesIO

import pytest
from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.models import Customer, Product, ProductCategory, Quotation, QuotationLine, Role, User
from app.models.enums import CustomerTier, QuotationStatus, RoleCode
from app.services.reports import (
    Report,
    ReportFilters,
    ReportRow,
    build_report,
    render_pdf,
    render_xlsx,
)

NOW = datetime.now(UTC)
_counter = itertools.count()


def _row(quote_number: str, total: str, discount: str) -> ReportRow:
    return ReportRow(
        quote_number=quote_number,
        customer_name="Acme Corp",
        owner_name="Test Rep",
        status="sent",
        total_amount=Decimal(total),
        discount_amount=Decimal(discount),
        blended_risk_score=Decimal("1.33"),
        created_at=NOW,
    )


# ---------------------------------------------------------------------------
# Pure rendering
# ---------------------------------------------------------------------------


def test_render_xlsx_is_a_valid_workbook_with_a_header_and_total_row() -> None:
    report = Report(
        rows=(_row("Q-1001", "1000.00", "100.00"), _row("Q-1002", "2000.00", "200.00")),
        total_amount=Decimal("3000.00"),
        total_discount=Decimal("300.00"),
    )

    content = render_xlsx(report)
    wb = load_workbook(BytesIO(content))
    ws = wb.active

    assert ws.cell(row=1, column=1).value == "Quote #"
    assert ws.cell(row=2, column=1).value == "Q-1001"
    assert ws.cell(row=3, column=1).value == "Q-1002"
    assert ws.cell(row=4, column=1).value == "TOTAL"
    assert ws.cell(row=4, column=5).value == "3000.00"


def test_render_xlsx_of_an_empty_report_still_has_headers_and_total() -> None:
    report = Report(rows=(), total_amount=Decimal("0"), total_discount=Decimal("0"))

    content = render_xlsx(report)
    wb = load_workbook(BytesIO(content))
    ws = wb.active

    assert ws.cell(row=1, column=1).value == "Quote #"
    assert ws.cell(row=2, column=1).value == "TOTAL"


def test_render_pdf_produces_non_empty_bytes_starting_with_the_pdf_magic_number() -> None:
    report = Report(
        rows=(_row("Q-1001", "1000.00", "100.00"),),
        total_amount=Decimal("1000.00"),
        total_discount=Decimal("100.00"),
    )

    content = render_pdf(report)

    assert len(content) > 0
    assert content[:4] == b"%PDF"


# ---------------------------------------------------------------------------
# Database orchestration
# ---------------------------------------------------------------------------


@pytest.fixture
async def rep_role(db_session: AsyncSession) -> Role:
    existing = (
        await db_session.execute(select(Role).where(Role.code == RoleCode.SALES_REP))
    ).scalar_one_or_none()
    if existing is not None:
        return existing
    role = Role(code=RoleCode.SALES_REP, name="Sales Rep")
    db_session.add(role)
    await db_session.flush()
    return role


@pytest.fixture
async def rep(db_session: AsyncSession, rep_role: Role) -> User:
    user = User(
        email=f"report-rep-{NOW.timestamp()}@example.test",
        password_hash="not-a-real-hash",
        full_name="Test Rep",
        role_id=rep_role.id,
    )
    db_session.add(user)
    await db_session.flush()
    return user


@pytest.fixture
async def customer(db_session: AsyncSession) -> Customer:
    customer = Customer(code=f"CUST-{NOW.timestamp()}", name="Acme Corp", tier=CustomerTier.GOLD)
    db_session.add(customer)
    await db_session.flush()
    return customer


@pytest.fixture
async def product(db_session: AsyncSession) -> Product:
    category = ProductCategory(code=f"CAT-{NOW.timestamp()}", name="Hardware")
    db_session.add(category)
    await db_session.flush()
    product = Product(
        sku=f"SKU-{NOW.timestamp()}",
        name="Test Laptop",
        category_id=category.id,
        list_price=Decimal("1000.00"),
        cost_price=Decimal("700.00"),
    )
    db_session.add(product)
    await db_session.flush()
    return product


async def _quotation(
    db_session: AsyncSession,
    customer: Customer,
    rep: User,
    product: Product,
    *,
    status,
    discount: str,
) -> Quotation:
    quotation = Quotation(
        quote_number=f"Q-REPORT-{NOW.timestamp()}-{next(_counter)}",
        customer_id=customer.id,
        owner_id=rep.id,
        status=status,
        subtotal_amount=Decimal("1000.00"),
        discount_amount=Decimal(discount),
        total_amount=Decimal("1000.00") - Decimal(discount),
    )
    db_session.add(quotation)
    await db_session.flush()
    db_session.add(
        QuotationLine(
            quotation_id=quotation.id,
            line_number=1,
            product_id=product.id,
            quantity=Decimal("1"),
            unit_list_price=product.list_price,
            unit_cost_price=product.cost_price,
        )
    )
    await db_session.flush()
    return quotation


async def test_report_scoped_to_one_owner_excludes_others(
    db_session: AsyncSession, customer: Customer, rep: User, product: Product
) -> None:
    await _quotation(db_session, customer, rep, product, status=QuotationStatus.SENT, discount="0")
    other_role = (await db_session.execute(select(Role))).scalars().first()
    other_rep = User(
        email=f"other-rep-{NOW.timestamp()}@example.test",
        password_hash="x",
        full_name="Other Rep",
        role_id=other_role.id,
    )
    db_session.add(other_rep)
    await db_session.flush()
    await _quotation(
        db_session, customer, other_rep, product, status=QuotationStatus.SENT, discount="0"
    )

    report = await build_report(db_session, ReportFilters(), owner_scope=rep.id)

    assert len(report.rows) == 1
    assert report.rows[0].owner_name == "Test Rep"


async def test_report_filters_by_status(
    db_session: AsyncSession, customer: Customer, rep: User, product: Product
) -> None:
    await _quotation(db_session, customer, rep, product, status=QuotationStatus.DRAFT, discount="0")
    await _quotation(db_session, customer, rep, product, status=QuotationStatus.SENT, discount="0")

    report = await build_report(db_session, ReportFilters(status="sent"), owner_scope=rep.id)

    assert len(report.rows) == 1
    assert report.rows[0].status == "sent"


async def test_report_totals_sum_correctly(
    db_session: AsyncSession, customer: Customer, rep: User, product: Product
) -> None:
    await _quotation(
        db_session, customer, rep, product, status=QuotationStatus.SENT, discount="100.00"
    )
    await _quotation(
        db_session, customer, rep, product, status=QuotationStatus.SENT, discount="200.00"
    )

    report = await build_report(db_session, ReportFilters(), owner_scope=rep.id)

    assert report.total_discount == Decimal("300.00")
    assert len(report.rows) == 2
